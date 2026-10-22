"""DOCX report export. Built in memory and returned as bytes (nothing is saved)."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .compare import BAND_COLOURS, GREY, format_value

# Shown at the top and bottom of every report, and in the app.
DISCLAIMER = (
    "This report does not indicate whether AI was used. Differences in writing can have many "
    "legitimate causes, including time pressure, topic, language background, editing tools such "
    "as spell-checkers, and accessibility accommodations. Use this information only to guide a "
    "structured follow-up conversation. The final judgement rests with the practitioner."
)

PROVISIONAL_NOTE = (
    "The reference ranges used for banding are provisional defaults and have not yet been "
    "calibrated on genuine candidate writing. Treat the bands as indicative only."
)


def _shade(cell, hex_colour: str) -> None:
    """Fill a Word table cell with a background colour (python-docx has no direct API)."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_colour.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shading)


def _disclaimer(doc: Document) -> None:
    """Add the disclaimer as a bold, dark red paragraph."""
    para = doc.add_paragraph()
    run = para.add_run(DISCLAIMER)
    run.bold = True
    run.font.color.rgb = RGBColor(0x8A, 0x1C, 0x1C)


def build_report(result: dict, paragraphs: list[dict], questions: list[str],
                 question_mode: str, notes: str, candidate_ref: str = "") -> bytes:
    """Build the full consistency report as a .docx file and return its bytes."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # --- title, date and disclaimer
    doc.add_heading("Written Response Consistency Report", level=0)
    meta = doc.add_paragraph()
    meta.add_run("Generated: ").bold = True
    meta.add_run(datetime.now().strftime("%Y-%m-%d %H:%M"))
    if candidate_ref:
        meta.add_run("\nCandidate reference: ").bold = True
        meta.add_run(candidate_ref)
    _disclaimer(doc)
    if result.get("provisional"):
        doc.add_paragraph(PROVISIONAL_NOTE).runs[0].italic = True

    # --- summary and length checks
    doc.add_heading("Summary", level=1)
    base, remote = result["baseline_features"], result["remote_features"]
    doc.add_paragraph(
        f"Baseline: {base['word_count']} words, {base['paragraph_count']} paragraph(s). "
        f"Remote: {remote['word_count']} words, {remote['paragraph_count']} paragraph(s)."
    )
    if result.get("length_message"):
        doc.add_paragraph(result["length_message"]).runs[0].italic = True
    if result["length_status"] == "refuse":
        doc.add_paragraph("No comparison was made.")
    else:
        doc.add_paragraph(result["summary"]).runs[0].bold = True

        # --- feature table, with the band cell shaded in its band colour
        doc.add_heading("Feature comparison", level=1)
        doc.add_paragraph(
            "Each difference is compared with the typical difference between two texts written "
            "by the same person. Green: within the typical range. Amber: larger than typical. "
            "Red: well beyond the typical range."
        )
        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Grid Accent 1"
        for cell, title in zip(table.rows[0].cells,
                               ["Feature", "Baseline", "Remote", "Difference", "Typical / high", "Band"]):
            cell.text = title
        for row in result["rows"]:
            d = row["decimals"]
            cells = table.add_row().cells
            cells[0].text = f"{row['label']} ({row['unit']})"
            cells[1].text = format_value(row["baseline"], d)
            cells[2].text = format_value(row["remote"], d)
            cells[3].text = format_value(row["difference"], d)
            cells[4].text = f"{format_value(row['typical'], d)} / {format_value(row['high'], d)}"
            cells[5].text = row["band"]
            _shade(cells[5], BAND_COLOURS[row["band"]])

        doc.add_heading("What the features mean", level=2)
        for row in result["rows"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(row["label"] + ": ").bold = True
            p.add_run(row["explanation"])

        # --- each remote paragraph in a shaded single-cell table
        doc.add_heading("Paragraph highlights (remote text)", level=1)
        doc.add_paragraph(
            "Paragraphs of 40+ words are ranked by how different they are from the baseline "
            "profile. Scores are relative and noisy at paragraph length; use them to choose "
            "what to ask about, not to draw conclusions."
        )
        for p in paragraphs:
            heading = doc.add_paragraph()
            label = f"Paragraph {p['index']} ({p['word_count']} words)"
            if p["eligible"]:
                label += f": score {p['score']:.2f}, {p['band'].lower()} (rank {p['rank']})"
                if p["drivers"]:
                    label += f"; largest differences: {', '.join(p['drivers'])}"
            else:
                label += ": too short to score"
            heading.add_run(label).bold = True
            table = doc.add_table(rows=1, cols=1)
            table.rows[0].cells[0].text = p["text"]
            _shade(table.rows[0].cells[0], BAND_COLOURS[p["band"] if p["eligible"] else GREY])

    # --- questions, notes and a closing disclaimer
    doc.add_heading("Suggested follow-up questions", level=1)
    doc.add_paragraph(f"Source: {question_mode}.").runs[0].italic = True
    for q in questions:
        doc.add_paragraph(q, style="List Number")

    doc.add_heading("Practitioner notes", level=1)
    doc.add_paragraph(notes.strip() or "(No notes recorded.)")

    doc.add_heading("Important", level=1)
    _disclaimer(doc)

    # Save to memory only, so no copy of the report is left on the server.
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
