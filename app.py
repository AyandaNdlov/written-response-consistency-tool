"""Written Response Consistency Tool: Streamlit UI.

Compares a supervised baseline writing sample with a remote (unproctored)
written assessment from the same candidate. All text is processed in memory
only; nothing is written to disk (POPIA).
"""

from __future__ import annotations

import html
from io import BytesIO
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import streamlit as st

from src.compare import (BAND_COLOURS, BAND_ICONS, GREY, compare_texts, format_value,
                         load_reference_ranges)
from src.paragraphs import analyse_paragraphs, flagged_paragraphs
from src.questions import LLM_MODEL, generate_questions, get_api_key
from src.report import DISCLAIMER, PROVISIONAL_NOTE, build_report

SAMPLES_DIR = Path(__file__).parent / "data" / "samples"
EXAMPLES = {
    "Genuine pair (same person, both genuine)": ("baseline_person1.txt", "remote_genuine_person1.txt"),
    "Genuine baseline vs AI-generated remote": ("baseline_person1.txt", "remote_ai_person1.txt"),
    "Genuine baseline vs AI-generated, partly edited": ("baseline_person1.txt", "remote_ai_edited_person1.txt"),
}

# Fixed display ranges for the radar chart, so both texts share one scale.
RADAR_FEATURES = {
    "sentence_len_mean": ("Sentence length", 5, 35),
    "sentence_len_sd": ("Sentence variation", 0, 15),
    "avg_word_len": ("Word length", 3.5, 6.0),
    "mtld": ("Vocab diversity", 30, 150),
    "long_word_pct": ("Long words %", 5, 40),
    "commas_p100": ("Commas", 0, 10),
    "discourse_markers_p100": ("Stock connectives", 0, 4),
    "contractions_p100": ("Contractions", 0, 6),
    "first_person_p100": ("First person", 0, 12),
    "flesch_reading_ease": ("Reading ease", 0, 100),
}

# Palette shared by the CSS, the hero illustration and the chart.
INK = "#1E3A2F"
CREAM = "#F7F4EE"
PINK = "#E8508A"
PINK_SOFT = "#FFD6E5"
PERIWINKLE = "#CFDBFF"
LIME = "#EEF57E"

st.set_page_config(page_title="Written Response Consistency Tool", page_icon="📝", layout="wide")


# ---------------------------------------------------------------- styling

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Gochi+Hand&display=swap');

:root {
  --ink: #1E3A2F; --cream: #F7F4EE; --pink: #E8508A; --pink-soft: #FFD6E5;
  --peri: #CFDBFF; --lime: #EEF57E; --hand: 'Gochi Hand', cursive;
}
[data-testid="stHeader"] { background: transparent; }
.block-container { padding-top: 1.2rem; max-width: 1240px; }

/* top bar */
.topbar { display: flex; align-items: center; justify-content: space-between; gap: 1rem;
  flex-wrap: wrap; padding: 0.4rem 0 1.2rem; }
.brand { display: flex; align-items: center; gap: 0.6rem; font-family: 'Bricolage Grotesque', sans-serif;
  font-weight: 800; font-size: 1.45rem; color: var(--ink); letter-spacing: -0.02em; }
.brand-dot { color: var(--pink); }
.nav { display: flex; gap: 1.6rem; flex-wrap: wrap; }
.nav a { color: var(--ink) !important; text-decoration: none !important; font-weight: 500; }
.nav a:hover { color: var(--pink) !important; }
.pill { background: var(--peri); color: var(--ink); border-radius: 999px; padding: 0.65rem 1.3rem;
  font-weight: 600; white-space: nowrap; }

/* hero */
.hero { text-align: center; padding: 2.2rem 0 0.6rem; }
.hero h1 { font-family: 'Bricolage Grotesque', sans-serif !important; font-weight: 800 !important;
  font-size: clamp(2.3rem, 5.6vw, 4.6rem) !important; line-height: 1.03 !important;
  letter-spacing: -0.035em; color: var(--ink) !important; max-width: 1000px; margin: 0 auto !important;
  padding: 0 !important; }
.circled { position: relative; display: inline-block; white-space: nowrap; }
.circled svg { position: absolute; left: -9%; top: -14%; width: 118%; height: 132%; pointer-events: none; }
.marked { background: linear-gradient(180deg, transparent 62%, var(--lime) 62%, var(--lime) 92%, transparent 92%);
  padding: 0 0.08em; }
.hero p.lede { font-size: 1.2rem; max-width: 640px; margin: 1.3rem auto 0; opacity: 0.85; }
.hand { font-family: var(--hand); font-size: 1.35rem; color: var(--ink); }
.stripe { display: flex; height: 14px; margin: 2rem 0 0; }
.stripe span:first-child { flex: 1; background: var(--ink); }
.stripe span:last-child { width: 9%; background: var(--pink); }
.doodle { display: block; width: 100%; max-width: 900px; height: auto; margin: 0 auto; }

/* notices */
.notice { display: flex; gap: 0.9rem; align-items: flex-start; background: #fff; border: 2px solid var(--ink);
  border-radius: 22px; padding: 1rem 1.3rem; margin: 1.4rem 0 1.6rem; box-shadow: 5px 5px 0 var(--pink); }
.notice .tag { background: var(--pink-soft); border-radius: 999px; padding: 0.15rem 0.75rem;
  font-family: var(--hand); font-size: 1.15rem; white-space: nowrap; }

/* section kickers */
.kicker { font-family: var(--hand); font-size: 1.5rem; color: var(--pink); margin: 2.2rem 0 -0.6rem;
  transform: rotate(-2deg); display: inline-block; }
.kicker .num { display: inline-flex; align-items: center; justify-content: center; width: 1.8rem;
  height: 1.8rem; border-radius: 50%; background: var(--ink); color: var(--cream); font-size: 1.05rem;
  margin-right: 0.4rem; }

/* cards (bordered containers) */
[data-testid="stVerticalBlockBorderWrapper"]:has(> div > [data-testid="stVerticalBlock"]) {
  background: #fff; }
div[data-testid="stTextArea"] textarea { font-size: 0.98rem; line-height: 1.55; background: var(--cream); }

/* buttons: dark ink pills that blush pink on hover, like the reference */
.stButton > button[kind="primary"], .stDownloadButton > button[kind="primary"],
[data-testid="stBaseButton-primary"] {
  background: var(--ink) !important; color: var(--cream) !important; border: 2px solid var(--ink) !important;
  padding: 0.75rem 2.2rem !important; font-weight: 700 !important; font-size: 1.05rem !important;
  transition: transform .15s ease, box-shadow .15s ease, background .15s ease; }
.stButton > button[kind="primary"]:hover, .stDownloadButton > button[kind="primary"]:hover,
[data-testid="stBaseButton-primary"]:hover {
  background: var(--pink) !important; border-color: var(--pink) !important; color: #fff !important;
  transform: translateY(-2px); box-shadow: 0 6px 0 var(--ink); }

/* metrics */
[data-testid="stMetric"] { background: #fff; border: 2px solid var(--ink); border-radius: 22px;
  padding: 1rem 1.2rem; }
[data-testid="stColumn"]:first-child [data-testid="stMetric"] { background: var(--pink-soft); }
[data-testid="stMetricValue"] { font-family: 'Bricolage Grotesque', sans-serif; font-weight: 800; }

/* paragraphs */
.para { background: #fff; border: 2px solid var(--ink); border-left: 12px solid var(--band);
  border-radius: 18px; padding: 1rem 1.2rem; margin-bottom: 0.9rem; }
.para-head { display: flex; flex-wrap: wrap; gap: 0.5rem; align-items: center; margin-bottom: 0.5rem; }
.para-head b { font-family: 'Bricolage Grotesque', sans-serif; font-size: 1.1rem; }
.chip { background: var(--band); border-radius: 999px; padding: 0.15rem 0.7rem; font-size: 0.85rem;
  font-weight: 600; }
.para-meta { font-size: 0.85rem; opacity: 0.7; }
.para-text { line-height: 1.6; }

/* questions */
.q { display: flex; gap: 1rem; align-items: flex-start; background: #fff; border: 2px solid var(--ink);
  border-radius: 18px; padding: 0.9rem 1.2rem; margin-bottom: 0.7rem; }
.q:nth-child(odd) { background: var(--cream); }
.qn { flex: none; display: inline-flex; align-items: center; justify-content: center; width: 2.1rem;
  height: 2.1rem; border-radius: 50%; background: var(--pink); color: #fff; font-weight: 700; }

/* sidebar */
[data-testid="stSidebar"] h2, [data-testid="stSidebar"] h3 { color: var(--cream) !important; }
.side-hand { font-family: var(--hand); font-size: 1.25rem; color: #FF8FB8; }
/* paper notes taped into the sidebar */
.paper { position: relative; color: var(--ink); background: var(--cream); border-radius: 4px 4px 14px 4px;
  padding: 1.4rem 1rem 0.9rem; margin: 1.6rem 0.2rem 1.2rem; transform: rotate(-1.6deg);
  box-shadow: 4px 6px 0 rgba(0,0,0,0.25); font-size: 0.9rem; line-height: 1.4; }
.paper.sticky { background: var(--peri); transform: rotate(1.4deg); border-radius: 4px 14px 4px 4px; }
.paper::before { content: ""; position: absolute; top: -11px; left: 50%; width: 84px; height: 22px;
  transform: translateX(-50%) rotate(-3deg); background: rgba(255,143,184,0.85); border-radius: 2px; }
.paper.sticky::before { background: rgba(238,245,126,0.9); transform: translateX(-50%) rotate(4deg); }
.paper .title { font-family: var(--hand); font-size: 1.45rem; line-height: 1; margin-bottom: 0.7rem; }
.paper .title u { text-decoration-color: var(--pink); text-decoration-thickness: 3px; text-underline-offset: 4px; }
.paper ul { list-style: none; margin: 0; padding: 0; }
.paper li { display: flex; gap: 0.55rem; align-items: flex-start; margin: 0 0 0.55rem; padding: 0; }
.paper li::before { flex: none; width: 1.3rem; height: 1.3rem; border-radius: 50%; display: inline-flex;
  align-items: center; justify-content: center; font-size: 0.75rem; font-weight: 700; margin-top: 0.05rem; }
.paper li.yes::before { content: "✓"; background: var(--ink); color: var(--lime); }
.paper li.no::before { content: "✕"; background: var(--pink); color: #fff; }
.paper li.lock::before { content: "🔒"; font-size: 0.85rem; }
.paper .sign { font-family: var(--hand); font-size: 1.05rem; text-align: right; opacity: 0.75; margin-top: 0.2rem; }

</style>
"""

LOGO = f"""<svg width="40" height="40" viewBox="0 0 40 40" fill="none" aria-hidden="true">
<rect x="4" y="7" width="20" height="26" rx="4" transform="rotate(-8 14 20)" stroke="{INK}" stroke-width="3"/>
<rect x="15" y="6" width="20" height="26" rx="4" transform="rotate(7 25 19)" fill="{PINK_SOFT}" stroke="{INK}" stroke-width="3"/>
<path d="M20 19.5l3.2 3.2 6-6.5" stroke="{PINK}" stroke-width="3.2" stroke-linecap="round" stroke-linejoin="round"/>
</svg>"""

# Hand-drawn ellipse around the circled hero word.
SCRIBBLE = f"""<svg viewBox="0 0 300 100" preserveAspectRatio="none" aria-hidden="true">
<path d="M150 8C70 6 8 22 10 52c2 30 80 42 150 40 72-2 132-18 130-46C288 18 210 4 120 12"
 fill="none" stroke="{PINK}" stroke-width="5" stroke-linecap="round" opacity="0.85"/></svg>"""

# Line-art doodle: two pages compared, with handwritten labels.
DOODLE = f"""<svg class="doodle" viewBox="0 0 900 250" fill="none" stroke="{INK}" stroke-width="3.2"
 stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">
<path d="M70 200c60-40 110-60 170-70" stroke="{LIME}" stroke-width="22" opacity="0.9"/>
<g transform="rotate(-7 230 120)">
  <rect x="170" y="40" width="130" height="165" rx="12" fill="#fff"/>
  <path d="M192 75h86M192 98h70M192 121h86M192 144h52M192 167h78"/>
</g>
<g transform="rotate(6 670 120)">
  <rect x="600" y="35" width="130" height="165" rx="12" fill="{PINK_SOFT}"/>
  <path d="M622 70h86M622 93h86M622 116h60M622 139h86M622 162h40"/>
</g>
<path d="M330 120c40-30 90-36 120-10s80 26 120-8" stroke-dasharray="2 12"/>
<path d="M556 92l16 10-14 14"/>
<circle cx="450" cy="92" r="34" fill="#fff"/>
<path d="M474 116l26 26" stroke-width="7"/>
<path d="M437 84c0-12 26-12 26 2 0 10-13 10-13 20M450 118v1" stroke="{PINK}" stroke-width="4"/>
<path d="M790 60c30 10 40 50 10 70s-10 50 30 40" stroke="{PERIWINKLE}" stroke-width="8"/>
<path d="M95 70l12 12M107 70l-12 12M780 200l12 12M792 200l-12 12" stroke="{PINK}"/>
<text x="150" y="236" font-family="Gochi Hand, cursive" font-size="26" fill="{INK}" stroke="none">baseline</text>
<text x="620" y="236" font-family="Gochi Hand, cursive" font-size="26" fill="{INK}" stroke="none">remote</text>
<text x="395" y="190" font-family="Gochi Hand, cursive" font-size="26" fill="{PINK}" stroke="none">same voice?</text>
</svg>"""

# Inject the custom styles once, before anything else is drawn.
st.markdown(CSS, unsafe_allow_html=True)


# ---------------------------------------------------------------- helpers

def read_upload(uploaded) -> str:
    """Read an uploaded .txt or .docx file into plain text, in memory."""
    data = uploaded.getvalue()
    if uploaded.name.lower().endswith(".docx"):
        from docx import Document
        return "\n\n".join(p.text for p in Document(BytesIO(data)).paragraphs if p.text.strip())
    return data.decode("utf-8", errors="replace")


def on_upload(which: str) -> None:
    """Copy an uploaded file into the matching text box and clear any old results."""
    uploaded = st.session_state.get(f"{which}_upload")
    if uploaded is not None:
        st.session_state[f"{which}_text"] = read_upload(uploaded)
        st.session_state.pop("analysis", None)


def on_example() -> None:
    """Fill both text boxes with the chosen example pair."""
    choice = st.session_state.get("example")
    if choice in EXAMPLES:
        base_file, remote_file = EXAMPLES[choice]
        st.session_state["baseline_text"] = (SAMPLES_DIR / base_file).read_text(encoding="utf-8")
        st.session_state["remote_text"] = (SAMPLES_DIR / remote_file).read_text(encoding="utf-8")
        st.session_state.pop("analysis", None)


def radar_chart(base: dict, remote: dict) -> go.Figure:
    """Radar chart of both texts, green for the baseline and pink for the remote."""
    labels, b_vals, r_vals = [], [], []
    for key, (label, lo, hi) in RADAR_FEATURES.items():
        labels.append(label)
        # Scale each value to 0-1 on its fixed range, clipping anything outside it.
        b_vals.append(min(max((base[key] - lo) / (hi - lo), 0), 1))
        r_vals.append(min(max((remote[key] - lo) / (hi - lo), 0), 1))
    fig = go.Figure()
    # Repeat the first point at the end so each shape closes.
    for name, vals, colour, fill in (("Baseline", b_vals, INK, "rgba(30,58,47,0.18)"),
                                     ("Remote", r_vals, PINK, "rgba(232,80,138,0.25)")):
        fig.add_trace(go.Scatterpolar(r=vals + vals[:1], theta=labels + labels[:1], name=name,
                                      fill="toself", fillcolor=fill,
                                      line=dict(color=colour, width=3), marker=dict(size=6)))
    fig.update_layout(polar=dict(bgcolor="rgba(0,0,0,0)",
                                 radialaxis=dict(visible=True, range=[0, 1], showticklabels=False,
                                                 gridcolor="rgba(30,58,47,0.15)"),
                                 angularaxis=dict(gridcolor="rgba(30,58,47,0.15)",
                                                  tickfont=dict(size=13))),
                      paper_bgcolor="rgba(0,0,0,0)", font=dict(family="DM Sans, sans-serif", color=INK),
                      margin=dict(l=50, r=50, t=30, b=30), height=460,
                      legend=dict(orientation="h", y=-0.08, x=0.5, xanchor="center", font=dict(size=14)))
    return fig


def band_style(val: str) -> str:
    """Table cell style for the Band column, e.g. "🟢 Within typical range"."""
    colour = BAND_COLOURS.get(val.split(" ", 1)[-1])
    return f"background-color: {colour}; color: {INK}; font-weight: 600" if colour else ""


def kicker(text: str, num: int | None = None) -> None:
    """Small handwritten label above a section, with an optional numbered badge."""
    badge = f"<span class='num'>{num}</span>" if num else ""
    st.markdown(f"<div class='kicker'>{badge}{text}</div>", unsafe_allow_html=True)


# ---------------------------------------------------------------- sidebar

api_key = get_api_key(st.secrets)
ranges = load_reference_ranges()

with st.sidebar:
    st.markdown("<div class='side-hand'>hey, practitioner!</div>", unsafe_allow_html=True)
    st.header("About this tool")
    st.markdown(
        "Compares a candidate's **supervised baseline** writing sample with their **remote** "
        "written assessment. It shows where writing patterns are similar or different, "
        "measured against how much one person's writing *typically* varies, and suggests "
        "follow-up interview questions."
    )
    st.markdown(
        "<div class='paper'><div class='title'>the <u>ground rules</u></div><ul>"
        "<li class='no'><span><b>Not an AI detector.</b> No AI score, no probability.</span></li>"
        "<li class='no'><span>Never the <b>only</b> basis for a selection decision.</span></li>"
        "<li class='yes'><span>Use it to plan a <b>fair, structured</b> follow-up chat.</span></li>"
        "<li class='yes'><span>Tell candidates how their written work is reviewed.</span></li>"
        "</ul><div class='sign'>~ you make the call</div></div>"
        "<div class='paper sticky'><div class='title'>privacy, promised</div><ul>"
        "<li class='lock'><span>Text stays <b>in memory only</b>. Nothing is saved (POPIA).</span></li>"
        "<li class='lock'><span>Remove names and ID numbers before pasting.</span></li>"
        "</ul></div>",
        unsafe_allow_html=True)
    # The Claude option only appears when an API key is set; otherwise template questions are used.
    use_llm = False
    if api_key:
        st.subheader("Question generation")
        use_llm = st.toggle("Use Claude to tailor questions", value=False,
                            help=f"Uses {LLM_MODEL}.")
        if use_llm:
            st.caption("When on, the flagged paragraphs are sent to Anthropic's API to write questions. "
                       "Make sure your organisation's data-processing agreement allows this.")
    if ranges.get("_meta", {}).get("status") != "calibrated":
        st.caption("⚠️ Reference ranges are **provisional** until calibrated (see README).")


# ---------------------------------------------------------------- hero

st.markdown(
    f"""<div class="topbar">
  <div class="brand">{LOGO}<span>Consistency<span class="brand-dot">.</span>Check</span></div>
  <nav class="nav"><a href="#compare">Compare</a><a href="#consistency-report">Report</a>
    <a href="#suggested-follow-up-questions">Questions</a></nav>
  <span class="pill">Not an AI detector</span>
</div>
<div class="hero">
  <h1>Compare writing <span class="circled">fairly{SCRIBBLE}</span>, then ask
  <span class="marked">better questions.</span></h1>
  <p class="lede">For assessment practitioners in graduate recruitment. Supports professional
  judgement; does not replace it.</p>
</div>
<div class="stripe"><span></span><span></span></div>
{DOODLE}""",
    unsafe_allow_html=True)

st.markdown(f"<div class='notice'><span class='tag'>please note</span><div>{html.escape(DISCLAIMER)}"
            "</div></div>", unsafe_allow_html=True)


# ---------------------------------------------------------------- inputs

kicker("pop in the two texts", 1)
st.header("Compare", anchor="compare")
st.selectbox("Load example", ["(none)"] + list(EXAMPLES), key="example", on_change=on_example)

col_b, col_r = st.columns(2, gap="large")
with col_b, st.container(border=True):
    st.subheader("Baseline text (supervised)")
    st.file_uploader("Upload .txt or .docx", type=["txt", "docx"], key="baseline_upload",
                     on_change=on_upload, args=("baseline",))
    st.text_area("Or paste text", key="baseline_text", height=300)
with col_r, st.container(border=True):
    st.subheader("Remote text (unproctored)")
    st.file_uploader("Upload .txt or .docx", type=["txt", "docx"], key="remote_upload",
                     on_change=on_upload, args=("remote",))
    st.text_area("Or paste text", key="remote_text", height=300)

st.write("")
# Run the analysis and keep it in session state, so typing notes later does not recompute it.
if st.button("Compare texts  →", type="primary"):
    baseline = st.session_state.get("baseline_text", "")
    remote = st.session_state.get("remote_text", "")
    if not baseline.strip() or not remote.strip():
        st.warning("Please provide both a baseline and a remote text.")
    else:
        with st.spinner("Analysing writing features..."):
            result = compare_texts(baseline, remote, ranges)
            paragraphs, questions, mode = [], [], "Template"
            if result["length_status"] != "refuse":
                paragraphs = analyse_paragraphs(remote, result["baseline_features"], ranges)
                questions, mode = generate_questions(flagged_paragraphs(paragraphs), paragraphs,
                                                     use_llm=use_llm, api_key=api_key)
        st.session_state["analysis"] = dict(result=result, paragraphs=paragraphs,
                                            questions=questions, mode=mode)


# ---------------------------------------------------------------- report

analysis = st.session_state.get("analysis")
if analysis:
    result = analysis["result"]
    st.markdown("<div class='stripe'><span></span><span></span></div>", unsafe_allow_html=True)
    kicker("here's what we found", 2)
    st.header("Consistency report", anchor="consistency-report")

    if result["length_status"] == "refuse":
        st.error(result["length_message"])
        st.stop()
    if result["length_status"] == "warn":
        st.warning(result["length_message"])
    if result["provisional"]:
        st.info(PROVISIONAL_NOTE)

    base, remote = result["baseline_features"], result["remote_features"]
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Features outside typical range", f"{result['n_outside']} of {result['n_assessed']}")
    m2.metric("🟠 Somewhat unusual", result["n_amber"])
    m3.metric("🔴 Notably different", result["n_red"])
    m4.metric("Words (baseline / remote)", f"{base['word_count']} / {remote['word_count']}")
    st.markdown(f"**Summary:** {result['summary']} This describes differences in writing "
                "patterns only. It is not a judgement about who wrote the text.")

    # --- feature table + radar side by side
    tab_col, radar_col = st.columns([1.25, 1], gap="large")
    with tab_col:
        st.subheader("Feature comparison")
        table = pd.DataFrame([{
            "Feature": r["label"],
            "Unit": r["unit"],
            "Baseline": format_value(r["baseline"], r["decimals"]),
            "Remote": format_value(r["remote"], r["decimals"]),
            "Difference": format_value(r["difference"], r["decimals"]),
            "Typical / high": f"{format_value(r['typical'], r['decimals'])} / {format_value(r['high'], r['decimals'])}",
            "Band": f"{BAND_ICONS[r['band']]} {r['band']}",
        } for r in result["rows"]])
        st.dataframe(table.style.map(band_style, subset=["Band"]), hide_index=True,
                     width="stretch", height=640)
    with radar_col:
        st.subheader("Writing profile")
        st.plotly_chart(radar_chart(base, remote), width="stretch")
        st.caption("Each axis is scaled to a fixed typical range (0 = low, 1 = high) so both texts "
                   "share one scale. Overlapping shapes indicate similar profiles.")
    with st.expander("What do these features mean?"):
        for r in result["rows"]:
            st.markdown(f"- **{r['label']}**: {r['explanation']}")
        st.markdown(f"- **Spelling evidence**: baseline has {base['british_count']} British / "
                    f"{base['american_count']} American variants; remote has {remote['british_count']} / "
                    f"{remote['american_count']}. Compared only when each text has at least 2.")
        if remote["discourse_markers_found"] or base["discourse_markers_found"]:
            st.markdown(f"- **Connectives found**: baseline {base['discourse_markers_found'] or 'none'}; "
                        f"remote {remote['discourse_markers_found'] or 'none'}")

    # --- paragraphs
    kicker("zooming in", 3)
    st.subheader("Paragraph highlights (remote text)")
    st.caption("Paragraphs of 40+ words are compared with the baseline and ranked. A stronger colour "
               "means more different from the baseline. Paragraph scores are noisy; use them to choose "
               "what to ask about, not to draw conclusions.")
    for p in analysis["paragraphs"]:
        colour = BAND_COLOURS[p["band"] if p["eligible"] else GREY]
        if p["eligible"]:
            chip = f"{BAND_ICONS[p['band']]} {p['band']}"
            meta = f"{p['word_count']} words · score {p['score']:.2f} · rank {p['rank']}"
            if p["drivers"]:
                meta += f" · largest differences: {', '.join(p['drivers'])}"
        else:
            chip, meta = "⚪ Not scored", f"{p['word_count']} words · fewer than 40 words"
        st.markdown(
            f"<div class='para' style='--band:{colour}'><div class='para-head'>"
            f"<b>Paragraph {p['index']}</b><span class='chip'>{chip}</span>"
            f"<span class='para-meta'>{html.escape(meta)}</span></div>"
            f"<div class='para-text'>{html.escape(p['text'])}</div></div>", unsafe_allow_html=True)

    # --- questions
    kicker("for the interview", 4)
    st.subheader("Suggested follow-up questions", anchor="suggested-follow-up-questions")
    st.caption(f"Source: {analysis['mode']}. Questions probe understanding and reasoning; edit them "
               "to fit your interview guide.")
    st.markdown("<div>" + "".join(
        f"<div class='q'><span class='qn'>{i}</span><div>{html.escape(q)}</div></div>"
        for i, q in enumerate(analysis["questions"], start=1)) + "</div>", unsafe_allow_html=True)

    # --- notes + report
    kicker("wrap it up", 5)
    st.subheader("Practitioner notes")
    with st.container(border=True):
        candidate_ref = st.text_input("Candidate reference (optional, avoid names)")
        notes = st.text_area("Notes (included in the report, not stored)", height=150)
        st.download_button(
            "Download report (.docx)  ↓",
            data=build_report(result, analysis["paragraphs"], analysis["questions"], analysis["mode"],
                              notes, candidate_ref),
            file_name="consistency_report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
        )
