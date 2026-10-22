from io import BytesIO
from pathlib import Path

import pytest

from src.compare import (AMBER, GREEN, GREY, RED, band_for, check_lengths, compare_texts,
                         load_reference_ranges)
from src.features import (extract_features, function_word_distance, function_word_profile, mtld,
                          spelling_counts, split_paragraphs, split_sentences, tokenize)
from src.paragraphs import analyse_paragraphs, flagged_paragraphs
from src.questions import generate_questions, key_quote
from src.report import DISCLAIMER, build_report

SAMPLES = Path(__file__).resolve().parents[1] / "data" / "samples"


def sample(name: str) -> str:
    return (SAMPLES / name).read_text(encoding="utf-8")


def test_tokenize_keeps_contractions_and_curly_apostrophes():
    assert tokenize("I don’t think it's 2 easy.") == ["I", "don't", "think", "it's", "easy"]


def test_sentence_split_handles_abbreviations():
    text = "We used tools, e.g. Excel. It worked! Did it scale? Mostly."
    assert split_sentences(text) == ["We used tools, e.g. Excel.", "It worked!", "Did it scale?", "Mostly."]


def test_paragraph_split_blank_lines_and_hard_wraps():
    assert len(split_paragraphs("One.\n\nTwo.\n\n\nThree.")) == 3
    # hard-wrapped lines that end mid-sentence stay one paragraph
    assert len(split_paragraphs("This line wraps\nin the middle of\na sentence.")) == 1


def test_mtld_higher_for_diverse_text():
    repetitive = ("the cat sat on the mat " * 20).split()
    diverse = ("every unique word appears here only once making the vocabulary rich varied "
               "interesting colourful distinct novel fresh lively bright clear").split()
    assert mtld(diverse) > mtld(repetitive)


def test_spelling_counts():
    assert spelling_counts("I will organise the programme and analyse colour behaviour.") == (5, 0)
    assert spelling_counts("We organized the program to analyze color and behavior.") == (0, 5)
    # 'analyses' (plural of analysis) and 'emphasis' are not counted
    assert spelling_counts("Two analyses with emphasis.") == (0, 0)


def test_extract_features_basic_counts():
    f = extract_features("I'd say it's fine. We can't stop now, can we?\n\nMoreover, it is crucial.")
    assert f["word_count"] == 14
    assert f["paragraph_count"] == 2
    assert f["sentence_count"] == 3
    assert f["contractions_p100"] == pytest.approx(300 / 14)
    assert f["discourse_markers_found"] == {"moreover": 1, "crucial": 1}


def test_function_word_distance_identity_and_range():
    prof = function_word_profile(tokenize(sample("baseline_person1.txt")))
    assert function_word_distance(prof, prof) == pytest.approx(0.0, abs=1e-9)
    other = function_word_profile(tokenize(sample("remote_ai_person1.txt")))
    assert 0 < function_word_distance(prof, other) <= 1


def test_banding():
    assert band_for(1.0, 2.0, 4.0) == GREEN
    assert band_for(3.0, 2.0, 4.0) == AMBER
    assert band_for(5.0, 2.0, 4.0) == RED
    assert band_for(None, 2.0, 4.0) == GREY


def test_length_checks():
    assert check_lengths(70, 300)[0] == "refuse"
    assert check_lengths(120, 300)[0] == "warn"
    assert check_lengths(200, 300)[0] == "ok"
    short = compare_texts("Too short. " * 10, sample("remote_ai_person1.txt"))
    assert short["length_status"] == "refuse" and short["rows"] == []


def test_reference_ranges_cover_every_feature():
    ranges = load_reference_ranges()
    result = compare_texts(sample("baseline_person1.txt"), sample("remote_genuine_person1.txt"), ranges)
    for row in result["rows"]:
        assert ranges[row["key"]]["high"] > ranges[row["key"]]["typical"] > 0


def test_genuine_pair_mostly_green_and_ai_pair_more_flagged():
    baseline = sample("baseline_person1.txt")
    genuine = compare_texts(baseline, sample("remote_genuine_person1.txt"))
    ai = compare_texts(baseline, sample("remote_ai_person1.txt"))
    edited = compare_texts(baseline, sample("remote_ai_edited_person1.txt"))
    assert genuine["n_outside"] <= 2 and genuine["n_red"] == 0
    assert ai["n_outside"] > genuine["n_outside"]
    assert ai["n_red"] >= 5
    assert genuine["n_outside"] <= edited["n_outside"] <= ai["n_outside"]


def test_paragraphs_and_template_questions():
    baseline = extract_features(sample("baseline_person1.txt"))
    remote = sample("remote_ai_edited_person1.txt")
    paras = analyse_paragraphs(remote, baseline, load_reference_ranges())
    flagged = flagged_paragraphs(paras)
    # the unedited generated paragraphs (2 and 3) should rank highest
    assert {p["index"] for p in flagged[:2]} == {2, 3}
    questions, mode = generate_questions(flagged, paras)
    assert mode == "Template"
    assert any("paragraph 2" in q.lower() for q in questions)
    banned = {"ai", "artificial", "chatgpt", "generated", "cheat", "cheating", "plagiarism"}
    for q in questions:
        assert not set(tokenize(q.lower())) & banned


def test_key_quote_truncates():
    assert key_quote("word " * 40).endswith("…")


def test_report_builds_and_contains_disclaimer():
    from docx import Document

    baseline, remote = sample("baseline_person1.txt"), sample("remote_ai_person1.txt")
    result = compare_texts(baseline, remote)
    paras = analyse_paragraphs(remote, result["baseline_features"], load_reference_ranges())
    data = build_report(result, paras, ["Q1?"], "Template", "Some notes", "CAND-01")
    text = "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)
    assert DISCLAIMER in text and "Some notes" in text
